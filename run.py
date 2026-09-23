"""Run all four questions and export reproducible JSON/CSV results."""

from __future__ import annotations

import argparse
import csv
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from figures import create_all
from coordination import solve as solve_coordination
from model import Scenario
from q1 import solve_all as solve_q1
from q2 import solve as solve_q2
from q3 import solve as solve_q3
from q4 import RESOURCE_KEYS, solve as solve_q4
from verify import verify_all, verify_q2


def save_json(path: Path, obj: object) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def save_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def transport_rows(rows: list[dict]) -> list[dict]:
    return [dict(sortie=r["id"], drone=r["drone"], model=r["model"], battery=r["battery"],
                 start_s=round(r["start"], 3), visits="→".join(z for z, _ in r["visits"]),
                 return_s=round(r["return_time"], 3), energy_kwh=round(r["energy_kwh"], 6),
                 return_soc_pct=round(r["return_soc"] * 100, 3),
                 boxes=";".join(b for _, boxes in r["visits"] for b in boxes))
            for r in rows]


def delivery_rows(deliveries: dict) -> list[dict]:
    return [dict(box=bid, sortie=x["sortie"], zone=x["zone"], delivered_s=round(x["time"], 3))
            for bid, x in sorted(deliveries.items())]


def write_q2_workbook(path: Path, q2: dict) -> None:
    """Fill only Q2 sheets in a copy of the supplied result template."""
    template = Path(__file__).resolve().parent.parent / "结果提交模板.xlsx"
    workbook = load_workbook(template)
    rows_by_sheet = {
        "Q2_运输架次": [
            (r["id"], r["drone"], r["model"], r["battery"], round(r["start"], 6),
             "→".join(zone for zone, _ in r["visits"]), round(r["return_time"], 6),
             round(r["energy_kwh"], 6))
            for r in q2["sorties"]
        ],
        "Q2_逐箱交付": [
            (bid, item["sortie"], item["zone"], round(item["time"], 6))
            for bid, item in sorted(q2["deliveries"].items())
        ],
    }
    for sheet_name, rows in rows_by_sheet.items():
        sheet = workbook[sheet_name]
        for row_number, values in enumerate(rows, 2):
            for column, value in enumerate(values, 1):
                target = sheet.cell(row_number, column, value)
                if row_number > 2:
                    source = sheet.cell(2, column)
                    target._style = copy(source._style)
                    target.alignment = copy(source.alignment)
                    target.protection = copy(source.protection)
    workbook.save(path)


def export_q2(output: Path, s: Scenario, q2: dict, checked: dict) -> None:
    output.mkdir(parents=True, exist_ok=True)
    save_json(output / "q2.json", q2)
    save_json(output / "q2_validation.json", checked)
    transport_fields = ["sortie", "drone", "model", "battery", "start_s", "visits", "return_s",
                        "energy_kwh", "return_soc_pct", "boxes"]
    delivery_fields = ["box", "sortie", "zone", "delivered_s"]
    save_csv(output / "q2_transport_sorties.csv", transport_fields, transport_rows(q2["sorties"]))
    save_csv(output / "q2_box_deliveries.csv", delivery_fields, delivery_rows(q2["deliveries"]))
    plans = {"soft_delay_priority": q2, **q2.get("alternatives", {})}
    tradeoff = []
    resource_rows = []
    deadline_rows = []
    for name, alternative in plans.items():
        tradeoff.append(dict(scheme=name, **alternative["objective"]))
        for row in alternative["sorties"]:
            common = dict(scheme=name, sortie=row["id"], model=row["model"])
            resource_rows.extend((
                dict(**common, resource_type="aircraft", resource_id=row["drone"],
                     activity="flight", start_s=row["start"], end_s=row["return_time"]),
                dict(**common, resource_type="battery", resource_id=row["battery"],
                     activity="flight", start_s=row["start"], end_s=row["return_time"]),
                dict(**common, resource_type="battery", resource_id=row["battery"],
                     activity="charging", start_s=row["return_time"], end_s=row["battery_ready"]),
            ))
        for bid, item in sorted(alternative["deliveries"].items()):
            box = s.boxes[bid]
            deadline_rows.append(dict(
                scheme=name, box=bid, kind=box.kind, priority=box.priority,
                desired_s=box.desired, hard_deadline_s=box.hard_deadline,
                delivered_s=item["time"], hard_slack_s=(box.hard_deadline - item["time"]
                                                         if box.hard_deadline is not None else ""),
                soft_delay_s=(max(0.0, item["time"] - box.desired)
                              if box.hard_deadline is None else "")))
    for name, alternative in q2.get("alternatives", {}).items():
        save_csv(output / f"q2_{name}_transport_sorties.csv", transport_fields,
                 transport_rows(alternative["sorties"]))
        save_csv(output / f"q2_{name}_box_deliveries.csv", delivery_fields,
                 delivery_rows(alternative["deliveries"]))
    save_csv(output / "q2_tradeoff.csv",
             ["scheme", "weighted_delivery_seconds", "weighted_soft_delay_seconds",
              "makespan", "energy_kwh", "sortie_count"], tradeoff)
    save_csv(output / "q2_resource_timeline.csv",
             ["scheme", "sortie", "model", "resource_type", "resource_id",
              "activity", "start_s", "end_s"], resource_rows)
    save_csv(output / "q2_deadline_audit.csv",
             ["scheme", "box", "kind", "priority", "desired_s", "hard_deadline_s",
              "delivered_s", "hard_slack_s", "soft_delay_s"], deadline_rows)
    write_q2_workbook(output / "q2_result_template.xlsx", q2)


def export(output: Path, s: Scenario, q1: dict, q2: dict, q3: dict, q4: dict, checks: dict) -> None:
    output.mkdir(parents=True, exist_ok=True)
    for name, obj in (("q1", q1), ("q3", q3), ("q4", q4), ("validation", checks)):
        save_json(output / f"{name}.json", obj)
    export_q2(output, s, q2, checks["q2"])

    payload = [dict(zone=z, model=g, safe_payload_kg=round(value, 6))
               for z, values in q1["max_safe_payload_kg"].items() for g, value in values.items()]
    save_csv(output / "q1_safe_payload.csv", ["zone", "model", "safe_payload_kg"], payload)
    batches = []
    for zone, plan in q1["zone_plans"].items():
        for i, b in enumerate(plan["batches"], 1):
            batches.append(dict(sortie=f"Q1-{zone}-{i:02d}", zone=zone, model=b["model"],
                                boxes=";".join(b["box_ids"]), kg=b["kg"], m3=b["volume"],
                                duration_s=round(b["duration"], 3), energy_kwh=round(b["energy"], 6),
                                return_soc_pct=round(b["return_soc"] * 100, 3)))
    save_csv(output / "q1_batches.csv",
             ["sortie", "zone", "model", "boxes", "kg", "m3", "duration_s", "energy_kwh", "return_soc_pct"],
             batches)
    if "reserve_sensitivity" in q1:
        sensitivity = [dict(reserve_fraction=rho, zone=z, model=g,
                            safe_payload_kg=round(value, 6), zone_sorties=data["by_zone"][z],
                            total_sorties=data["sorties"])
                       for rho, data in q1["reserve_sensitivity"].items()
                       for z, values in data["payload_kg"].items() for g, value in values.items()]
        save_csv(output / "q1_reserve_sensitivity.csv",
                 ["reserve_fraction", "zone", "model", "safe_payload_kg", "zone_sorties", "total_sorties"], sensitivity)

    transport_fields = ["sortie", "drone", "model", "battery", "start_s", "visits", "return_s",
                        "energy_kwh", "return_soc_pct", "boxes"]
    delivery_fields = ["box", "sortie", "zone", "delivered_s"]
    save_csv(output / "q3_transport_sorties.csv", transport_fields, transport_rows(q3["transport_sorties"]))
    save_csv(output / "q3_box_deliveries.csv", delivery_fields, delivery_rows(q3["deliveries"]))

    relay_rows = [dict(sortie=r["id"], relay=r["relay"], component=r["energy_component"],
                       start_s=round(r["start"], 3), lon=r["lon"], lat=r["lat"],
                       altitude_m=round(r["altitude"], 3), agl_m=r["agl"],
                       established_s=round(r["established"], 3), service_end_s=round(r["service_end"], 3),
                       return_s=round(r["return_time"], 3), energy_kwh=round(r["energy_kwh"], 6),
                       return_soc_pct=round(r["return_soc"] * 100, 3),
                       covered_transport=";".join(r["covered_transport"]))
                  for r in q3["relay_sorties"]]
    save_csv(output / "q3_relay_sorties.csv",
             ["sortie", "relay", "component", "start_s", "lon", "lat", "altitude_m", "agl_m",
              "established_s", "service_end_s", "return_s", "energy_kwh", "return_soc_pct", "covered_transport"],
             relay_rows)
    comm_rows = [dict(transport_sortie=c["sortie"], phase=c["phase"], start_s=round(c["start"], 3),
                      end_s=round(c["end"], 3), mode=c["mode"], relay_sortie=c["relay_sortie"] or "")
                 for c in q3["communication"] if c["end"] > c["start"] + 1e-8]
    save_csv(output / "q3_communication.csv",
             ["transport_sortie", "phase", "start_s", "end_s", "mode", "relay_sortie"], comm_rows)

    partition_rows = []
    for k, plan in q4["partitions"].items():
        if not plan["feasible"]:
            continue
        for scheme, variant in [("minimum_shortage", plan), *plan["alternatives"].items()]:
            for group in variant["groups"]:
                row = dict(k=k, scheme=scheme, group=group["id"], zones=";".join(group["zones"]),
                           boxes=group["boxes"], cargo_kg=group["cargo_kg"],
                           work_s=round(group["work_seconds"], 3),
                           energy_kwh=round(group["energy_kwh"], 6),
                           group_workload_cv=round(variant["workload_cv"], 6),
                           total_shortage=sum(variant["shortage"].values()) if scheme != "minimum_shortage"
                           else sum(variant["shortage_strict"].values()))
                row.update(group["strict_ids"])
                partition_rows.append(row)
    save_csv(output / "q4_partition_resources.csv",
             ["k", "scheme", "group", "zones", "boxes", "cargo_kg", "work_s", "energy_kwh",
              "group_workload_cv", "total_shortage", *RESOURCE_KEYS],
             partition_rows)

    summary = dict(input_counts=dict(service_areas=len(s.zone_boxes), boxes=len(s.boxes),
                                     transport_aircraft=len(s.aircraft), relay_aircraft=len(s.relays)),
                   q1=q1["totals"], q2=q2["objective"], q3=q3["objective"],
                   q3_validation=checks["q3"],
                   q4={k: dict(shortage=v["shortage_strict"],
                               resource_total=v["resource_total_strict"],
                               workload_cv=v["workload_cv"],
                               groups=[g["zones"] for g in v["groups"]],
                               alternatives={name: dict(total_shortage=a["total_shortage"],
                                                        workload_cv=a["workload_cv"],
                                                        groups=[g["zones"] for g in a["groups"]])
                                             for name, a in v["alternatives"].items()})
                       for k, v in q4["partitions"].items() if v["feasible"]})
    coordination = q3.get("search", {}).get("coordination")
    if coordination:
        summary["q3_coordination"] = {key: coordination[key] for key in (
            "baseline_objective", "selected_promotion", "candidates_evaluated",
            "feasible_candidates", "improvement_weighted_delivery_seconds")}
    save_json(output / "summary.json", summary)


def main() -> None:
    parser = argparse.ArgumentParser(description="Solve the D problem with the supplied data folder")
    default_output = Path(__file__).resolve().parent / "results"
    parser.add_argument("--output", type=Path, default=default_output)
    parser.add_argument("--no-sensitivity", action="store_true", help="Skip Q1 reserve sensitivity")
    parser.add_argument("--no-merge", action="store_true", help="Keep initial single-zone Q2 jobs")
    parser.add_argument("--only-q2", action="store_true", help="Solve and verify Q2, then fill its result-template sheets")
    parser.add_argument("--no-coordination", action="store_true", help="Use the original fixed-Q2 relay plan")
    parser.add_argument("--q3-seed", choices=("soft_delay_priority", "zero_delay_efficient",
                                              "delivery_priority", "energy_guarded",
                                              "energy_guarded_timely",
                                              "alns_zero_delay_efficient"),
                        default="energy_guarded", help="Q2 candidate used to initialize Q3")
    parser.add_argument("--coordination-rounds", type=int, default=2,
                        help="Communication-guided search rounds (default: 2)")
    parser.add_argument("--coordination-trials", type=int, default=3,
                        help="Full relay evaluations per round (default: 3)")
    parser.add_argument("--promotion-trials", type=int, default=12,
                        help="Soft-sortie insertion trials (default: 12)")
    parser.add_argument("--no-figures", action="store_true", help="Skip static result figures")
    args = parser.parse_args()
    scenario = Scenario()
    if args.only_q2:
        if args.output == default_output:
            args.output = Path(__file__).resolve().parent / "results_q2"
        print("Q2: transport and shared batteries", flush=True)
        q2 = solve_q2(scenario, improve=not args.no_merge)
        checked = verify_q2(scenario, q2)
        export_q2(args.output, scenario, q2, checked)
        print(json.dumps(dict(primary=q2["objective"],
                              alternatives={name: plan["objective"] for name, plan in q2["alternatives"].items()},
                              validation=checked), ensure_ascii=False, indent=2), flush=True)
        return
    print("Q1: exact per-zone partition", flush=True)
    q1 = solve_q1(scenario, sensitivity=not args.no_sensitivity)
    print("Q2: transport and shared batteries", flush=True)
    q2 = solve_q2(scenario, improve=not args.no_merge)
    print("Q3: communication-guided transport and relay search", flush=True)
    q3_seed = q2 if args.q3_seed == "soft_delay_priority" else q2["alternatives"][args.q3_seed]
    if args.no_coordination:
        q3 = solve_q3(scenario, q3_seed)
    else:
        q3 = solve_coordination(scenario, q3_seed, rounds=args.coordination_rounds,
                                trials_per_round=args.coordination_trials,
                                promotion_trials=args.promotion_trials)
    q3.setdefault("search", {})["q2_seed_policy"] = args.q3_seed
    print("Q4: independent task-group resources", flush=True)
    q4 = solve_q4(scenario, q3)
    print("Replay verification", flush=True)
    checks = verify_all(scenario, q1, q2, q3, q4)
    export(args.output, scenario, q1, q2, q3, q4, checks)
    if not args.no_figures:
        print("Drawing result figures", flush=True)
        create_all(scenario, q3, q4, args.output)
    print(f"Saved results to {args.output}", flush=True)
    print(json.dumps(dict(q1=q1["totals"], q2=q2["objective"], q3=q3["objective"],
                          q4_shortage={k: v["shortage_strict"] for k, v in q4["partitions"].items()}),
                     ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
